"""Refresh published observations from primary public sources; never interpolate.

No API keys, authentication, browser impersonation or bypass of HTTP blocks.
A failed source retains its previously verified series. A changed series must
also pass the production JavaScript validator before GitHub publishes it.
"""
import argparse
import copy
import datetime as dt
import io
import json
import os
import re
import unicodedata
import urllib.parse
import urllib.request
import urllib.error
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
CUB_INDEX = 'https://www.sinduscongoias.com.br/cub-custo-unitario-basico'
ANP_ROOT = 'https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/'
ANP_HIST = ANP_ROOT + 'precos-revenda-e-de-distribuicao-combustiveis/shlp/semanal/'
ANP_STATES = ANP_HIST + 'semanal-estados-desde-2013.xlsx'
ANP_CITIES = ANP_HIST + 'semanal-municipios-{year}.xlsx'
ANP_LATEST = ANP_ROOT + 'levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas'
BRENT = 'https://fred.stlouisfed.org/graph/fredgraph.csv?id=DCOILBRENTEU'
BRENT_EIA = 'https://www.eia.gov/dnav/pet/pet_pri_spt_s1_d.htm'
IPCA_MONTHLY = 'https://api.bcb.gov.br/dados/serie/bcdata.sgs.433/dados/ultimos/13?formato=json'
IPCA_HISTORY = 'https://api.bcb.gov.br/dados/serie/bcdata.sgs.433/dados?formato=json'
IPCA_12M = 'https://api.bcb.gov.br/dados/serie/bcdata.sgs.13522/dados/ultimos/13?formato=json'
IPCA_INDEX = 'https://apisidra.ibge.gov.br/values/t/1737/n1/all/v/2266/p/last%2013'
IPCA_REFERENCE = 'https://www.gov.br/fazenda/pt-br/central-de-conteudo/publicacoes/conjuntura-economica/inflacao/defeso-eleitoral-2026/informativo-ipca-set2026.html/view'
RECOVERIES = {}
MONTHS = ['janeiro','fevereiro','marco','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']


def normalized(value):
    text = unicodedata.normalize('NFKD', str(value or '')).translate(str.maketrans({c:'-' for c in '−‐‑‒–—'}))
    return re.sub(r'\s+', ' ', ''.join(c for c in text if not unicodedata.combining(c))).strip().upper()


def number(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace('−', '-').replace('%', '')
    if ',' in text:
        text = text.replace('.', '').replace(',', '.')
    return float(text)


def sgs_monthly(raw, source, maximum=1000):
    try:
        payload=json.loads(raw.decode('utf-8-sig') if isinstance(raw, bytes) else raw)
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError('BCB SGS response is not valid JSON') from error
    if not isinstance(payload, list) or not payload:
        raise ValueError('BCB SGS monthly response is empty')
    records={}
    for item in payload:
        if not isinstance(item, dict):
            raise ValueError('BCB SGS monthly record is invalid')
        stamp=str(item.get('data','')).strip()
        try:
            reference=dt.datetime.strptime(stamp, '%d/%m/%Y').date() if '/' in stamp else dt.date.fromisoformat(stamp[:10])
        except ValueError as error:
            raise ValueError('BCB SGS monthly date is invalid') from error
        value=number(item.get('valor'))
        if value <= -100 or value > maximum:
            raise ValueError('BCB SGS monthly value is invalid')
        period=reference.strftime('%Y-%m')
        if period in records:
            raise ValueError('BCB SGS monthly period is duplicated')
        records[period]={'period':period,'value':value,'sourceURL':source}
    return [records[period] for period in sorted(records)]


def sidra_index(raw, source, maximum=100000):
    """Parse an IBGE SIDRA monthly number-index response without interpolation."""
    try:
        payload=json.loads(raw.decode('utf-8-sig') if isinstance(raw, bytes) else raw)
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError('IBGE SIDRA response is not valid JSON') from error
    if not isinstance(payload, list) or not payload:
        raise ValueError('IBGE SIDRA response is empty')
    records={}
    for position, item in enumerate(payload):
        if position == 0 and isinstance(item, dict) and item.get('V') == 'Valor':
            # SIDRA includes a descriptive header, not an observation.
            continue
        if not isinstance(item, dict):
            raise ValueError('IBGE SIDRA record is invalid')
        value=item.get('V')
        if value in (None, '', '-', '..'):
            # SIDRA can return a metadata/header row; a missing observation is
            # not silently converted to zero.
            continue
        period=None
        for key, candidate in item.items():
            text=str(candidate).strip()
            if re.fullmatch(r'\d{6}', text) and 1 <= int(text[4:]) <= 12:
                period=text[:4]+'-'+text[4:]
                break
        if period is None:
            raise ValueError('IBGE SIDRA monthly period is invalid')
        indexed=number(value)
        if indexed <= 0 or indexed > maximum:
            raise ValueError('IBGE SIDRA index value is invalid')
        if period in records:
            raise ValueError('IBGE SIDRA monthly period is duplicated')
        records[period]={'period':period,'value':indexed,'sourceURL':source}
    if not records:
        raise ValueError('IBGE SIDRA response has no monthly observations')
    return [records[period] for period in sorted(records)]


def ipca_year_to_date(observations, period):
    year=period[:4]
    month=int(period[5:7])
    year_points=[item for item in observations if item['period'][:4]==year and int(item['period'][5:7])<=month]
    months=[int(item['period'][5:7]) for item in year_points]
    if months != list(range(1, month+1)):
        raise ValueError('IPCA monthly history has a gap before '+period)
    factor=1.0
    for item in year_points:
        factor*=1+item['value']/100
    return round((factor-1)*100, 2)


def refresh_ipca(previous, today):
    monthly=sgs_monthly(download(IPCA_MONTHLY), IPCA_MONTHLY, maximum=100)
    history=sgs_monthly(download(IPCA_HISTORY), IPCA_HISTORY, maximum=100)
    twelve_months=sgs_monthly(download(IPCA_12M), IPCA_12M, maximum=1000)
    index=sidra_index(download(IPCA_INDEX), IPCA_INDEX)
    current_month=today.strftime('%Y-%m')
    if not monthly or monthly[-1]['period']>current_month or not history or history[-1]['period']>current_month:
        raise ValueError('IPCA monthly response has a future reference')
    annual={item['period']:item for item in twelve_months}
    indexes={item['period']:item for item in index}
    observations=[]
    for item in monthly[-13:]:
        annual_item=annual.get(item['period'])
        if not annual_item:
            raise ValueError('IPCA 12-month response lacks '+item['period'])
        index_item=indexes.get(item['period'])
        if not index_item:
            raise ValueError('IBGE SIDRA index response lacks '+item['period'])
        record=copy.deepcopy(item)
        record['annualChange']=annual_item['value']
        record['annualChangeSourceURL']=IPCA_12M
        record['yearToDate']=ipca_year_to_date(history, item['period'])
        record['indexValue']=index_item['value']
        record['indexSourceURL']=IPCA_INDEX
        record['referenceURL']=IPCA_REFERENCE
        observations.append(record)
    if len(observations)<2:
        raise ValueError('IPCA monthly history is too short')
    result=copy.deepcopy(previous)
    result.update(sourceURL=IPCA_MONTHLY, historySourceURL=IPCA_HISTORY, annualSourceURL=IPCA_12M, indexSourceURL=IPCA_INDEX, referenceURL=IPCA_REFERENCE, annualSeries='BCB-SGS-13522', indexSeries='IBGE-SIDRA-1737-2266', observations=observations)
    return result


def day(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date() if isinstance(value, dt.datetime) else value
    text = str(value).strip()
    return dt.datetime.strptime(text[:10], '%d/%m/%Y').date() if '/' in text else dt.date.fromisoformat(text[:10])


def download(url):
    request = urllib.request.Request(url, headers={'User-Agent': 'Popular-Widget-Data/1.0'})
    for attempt in range(2):
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                if response.status != 200:
                    raise ValueError('HTTP response is not 200')
                raw = response.read(32*1024*1024+1)
            break
        except (TimeoutError, urllib.error.URLError) as error:
            reason = error.reason if isinstance(error, urllib.error.URLError) else error
            # One ordinary retry for a timeout only. Never retry HTTP blocks,
            # weaken TLS verification, impersonate a browser or switch proxies.
            if attempt or not isinstance(reason, TimeoutError):
                raise
    if len(raw) > 32*1024*1024:
        raise ValueError('Source exceeds ingestion size limit')
    return raw


class Page(HTMLParser):
    def __init__(self, raw):
        super().__init__()
        self.links, self.rows, self.text = [], [], []
        # Each table owns its row/cell state. EIA nests the product label in a
        # small table inside the outer row that owns the daily price columns.
        self.tables = [{'row': None, 'cell': None}]
        self.feed(raw.decode('utf-8-sig'))
        while self.tables:
            self.finish_row()
            self.tables.pop()

    def finish_cell(self):
        table = self.tables[-1]
        if table['cell'] is not None and table['row'] is not None:
            table['row'].append(' '.join(' '.join(table['cell']).split()))
        table['cell'] = None

    def finish_row(self):
        self.finish_cell()
        table = self.tables[-1]
        if table['row'] is not None:
            self.rows.append(table['row'])
        table['row'] = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'a' and attrs.get('href'):
            self.links.append(attrs['href'])
        if tag == 'table':
            self.tables.append({'row': None, 'cell': None})
        table = self.tables[-1]
        if tag == 'tr':
            self.finish_row()
            table['row'] = []
        if tag in ('td', 'th') and table['row'] is not None:
            # HTML permits omitting </td>/</th> before the next table cell.
            # HTMLParser emits tokens without browser-style implicit closure.
            self.finish_cell()
            table['cell'] = []

    def handle_data(self, data):
        self.text.append(data)
        for table in self.tables:
            if table['cell'] is not None:
                table['cell'].append(data)

    def handle_endtag(self, tag):
        if tag in ('td', 'th'):
            self.finish_cell()
        if tag in ('tr', 'table'):
            self.finish_row()
        if tag == 'table' and len(self.tables)>1:
            self.tables.pop()


def cub_text(text, period, source):
    clean = ' '.join(text.split())
    month = MONTHS[int(period[5:])-1]
    # PDF glyph extraction can separate letters (e.g. "MAR ÇO", "A L TO").
    month_pattern=r'\s*'.join(normalized(month))
    if not re.search(r'\b'+month_pattern+r'\s*-?\s*'+period[:4]+r'\b', normalized(clean)):
        raise ValueError('CUB reference month does not match publication')
    # Layout extraction keeps the three columns low / normal / high in order.
    # The separate "Valor referencial" and its changes describe R16-A and must
    # never be reused for the user's R1 normal series.
    if not re.search(r'PADRAO BAIXO\s+PADRAO NORMAL\s+PADRAO A\s*L\s*T\s*O', normalized(clean)):
        raise ValueError('CUB residential standard columns not found')
    values = re.findall(r'\bR\s*-?\s*1\s+(\d[\d.]*,\d{2})\b', clean, re.I)
    if len(values) != 3:
        raise ValueError('CUB R1 low/normal/high row is ambiguous')
    value = number(values[1])
    if value <= 0:
        raise ValueError('CUB R1-N cost invalid')
    return {'period': period, 'value': value, 'typology': 'R1-N', 'sourceURL': source}


def refresh_cub(previous, today):
    page = Page(download(CUB_INDEX))
    links = {}
    for href in page.links:
        url = urllib.parse.urljoin(CUB_INDEX, href)
        match = re.search(r'/cub-([a-z]+)-(\d{4})\.pdf$', url)
        if url.startswith('https://www.sinduscongoias.com.br/') and match and match[1] in MONTHS:
            period = f'{match[2]}-{MONTHS.index(match[1])+1:02}'
            if period <= today.strftime('%Y-%m'):
                links[period] = url
    if not links:
        raise ValueError('No official monthly CUB PDF links')
    records = {r['period']: r for r in previous['observations']}
    for period, url in sorted(links.items())[-14:]:
        raw = download(url)
        if not raw.startswith(b'%PDF'):
            raise ValueError('CUB response is not PDF')
        text = '\n'.join(p.extract_text(extraction_mode='layout') for p in PdfReader(io.BytesIO(raw)).pages)
        records[period] = cub_text(text, period, url)
    result = copy.deepcopy(previous)
    result['observations'] = sorted(records.values(), key=lambda r: r['period'])[-13:]
    result['sourceURL'] = result['observations'][-1]['sourceURL']
    result['regime'] = 'not-disclosed-in-source'
    return result


def sheet_records(raw, source):
    if not raw.startswith(b'PK'):
        raise ValueError('ANP response is not XLSX')
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        for sheet in book:
            header = None
            for cells in sheet.iter_rows(values_only=True):
                names = [normalized(c) for c in cells]
                if 'PRODUTO' in names and 'PRECO MEDIO REVENDA' in names and 'DATA FINAL' in names:
                    header = {n: i for i, n in enumerate(names) if n}
                    continue
                if header is None:
                    continue
                def get(*keys):
                    return next((cells[header[k]] for k in keys if k in header and header[k] < len(cells)), None)
                if normalized(get('PRODUTO')) != 'OLEO DIESEL S10':
                    continue
                if normalized(get('UNIDADE DE MEDIDA')) not in ('R$/L', 'R$/LITRO'):
                    raise ValueError('ANP diesel retail unit is not R$/L')
                if get('DATA FINAL') is None:
                    continue
                start, end = day(get('DATA INICIAL')), day(get('DATA FINAL'))
                state, city = normalized(get('ESTADO','ESTADOS','ESTADO - SIGLA')), normalized(get('MUNICIPIO'))
                value = number(get('PRECO MEDIO REVENDA'))
                if end-start != dt.timedelta(days=6) or value <= 0:
                    raise ValueError('Invalid ANP weekly retail observation')
                record = {'period': end.isoformat(), 'periodStart': start.isoformat(), 'value': value, 'sourceURL': source}
                for key, name in [('stations','NUMERO DE POSTOS PESQUISADOS'),('min','PRECO MINIMO REVENDA'),('max','PRECO MAXIMO REVENDA')]:
                    if get(name) is not None:
                        record[key] = int(number(get(name))) if key == 'stations' else number(get(name))
                yield state, city, record
    finally:
        book.close()


def refresh_diesel(previous, today):
    municipal_url = ANP_CITIES.format(year=today.year)
    municipal, coverage = {}, set()
    for state, city, record in sheet_records(download(municipal_url), municipal_url):
        coverage.add(record['period'])
        if state in ('GO','GOIAS') and city == 'JATAI':
            record['geography'] = 'JATAI-GO'
            municipal[record['period']] = record
    states = {}
    for state, city, record in sheet_records(download(ANP_STATES), ANP_STATES):
        if state in ('GO','GOIAS') and not city:
            record.update(geography='GO', noMunicipalResult=record['period'] in coverage and record['period'] not in municipal, municipalitySourceURL=municipal_url)
            states[record['period']] = record
    # The small latest-week workbook can be ahead of the historical download.
    page = Page(download(ANP_LATEST))
    releases = []
    for href in page.links:
        url = urllib.parse.urljoin(ANP_LATEST, href)
        match = re.search(r'resumo_semanal_lpc_(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})\.xlsx$', url)
        if url.startswith(ANP_ROOT) and match and match[2] <= today.isoformat():
            releases.append((match[2], url))
    if releases:
        period, url = max(releases)
        records = list(sheet_records(download(url), url))
        cities = [(s,c,r) for s,c,r in records if c]
        if not cities:
            raise ValueError('Latest ANP municipal coverage not found')
        own = [r for s,c,r in cities if s in ('GO','GOIAS') and c == 'JATAI']
        for r in own:
            r['geography'] = 'JATAI-GO'
            municipal[r['period']] = r
        for state, city, r in records:
            if state in ('GO','GOIAS') and not city:
                r.update(geography='GO', noMunicipalResult=not own, municipalitySourceURL=url)
                states[r['period']] = r
    if not states or not coverage:
        raise ValueError('ANP official state/municipal coverage absent')
    periods = sorted(set(states)|set(municipal))[-13:]
    observations = []
    for period in periods:
        if period in states:
            observations.append(states[period])
        if period in municipal:
            observations.append(municipal[period])
    result = copy.deepcopy(previous)
    result['observations'] = observations
    result['sourceURL'] = ANP_STATES
    return result


def refresh_brent_fred(previous, today):
    # Six months covers at least 90 ordinary business-day observations without
    # asking FRED to generate the complete history back to 1987 on each run.
    text = download(BRENT+'&cosd='+(today-dt.timedelta(days=180)).isoformat()).decode('utf-8-sig')
    lines = text.strip().splitlines()
    if not lines or lines.pop(0) not in ('observation_date,DCOILBRENTEU','DATE,DCOILBRENTEU'):
        raise ValueError('FRED Brent CSV schema invalid')
    observations = []
    for line in lines:
        period, value = line.split(',')
        if value in ('','.'):
            continue
        if day(period) > today or number(value) <= 0:
            raise ValueError('FRED future/invalid observation')
        observations.append({'period': period, 'value': number(value), 'sourceURL': 'https://fred.stlouisfed.org/series/DCOILBRENTEU'})
    result = copy.deepcopy(previous)
    result['observations'] = observations[-90:]
    result.pop('publishedAt', None)
    return result


def brent_eia_page(raw, previous, today):
    page = Page(raw)
    text = normalized(' '.join(page.text))
    if 'SPOT PRICES' not in text or 'CRUDE OIL IN DOLLARS PER BARREL' not in text:
        raise ValueError('EIA daily spot crude/USD per barrel header missing')
    date_rows = [[dt.datetime.strptime(c,'%m/%d/%Y' if len(c.split('/')[-1])==4 else '%m/%d/%y').date() for c in row if re.fullmatch(r'\d{1,2}/\d{1,2}/\d{2}(?:\d{2})?',c)] for row in page.rows]
    dates = next((row for row in date_rows if len(row)>=2), None)
    rows = [row for row in page.rows if any(re.fullmatch(r'BRENT\s*(?:-\s*)?EUROPE',normalized(c)) for c in row) and sum(bool(re.fullmatch(r'\d+\.\d+|\-|--|NA|W',c)) for c in row)>=2]
    if not dates or len(rows)!=1 or dates!=sorted(set(dates)):
        fragment=re.search(rb'.{0,40}Brent.{0,300}',raw,re.I|re.S)
        raise ValueError('EIA daily Brent-Europe/date row missing or ambiguous; public HTML near product='+(' '.join(fragment[0].decode('utf-8','replace').split()) if fragment else 'absent')+'; date rows='+str([r for r in page.rows if any(re.search(r"\d+/\d+/\d+",c) for c in r)][:2]))
    values = [c for c in rows[0] if re.fullmatch(r'\d+\.\d+|\-|--|NA|W',c)]
    if len(values)!=len(dates):
        raise ValueError(f'EIA Brent date/value columns disagree ({len(dates)} dates, {len(values)} values)')
    records = {r['period']:copy.deepcopy(r) for r in previous['observations']}
    received=[]
    for date, value in zip(dates, values):
        if value in ('-','--','NA','W'):
            continue
        if date>today or number(value)<=0:
            raise ValueError('EIA future/nonpositive Brent observation')
        records[date.isoformat()]={'period':date.isoformat(),'value':number(value),'sourceURL':BRENT_EIA}
        received.append(date.isoformat())
    if len(received)<2 or max(received)<previous['observations'][-1]['period']:
        raise ValueError('EIA Brent response missing current verified observations')
    result=copy.deepcopy(previous)
    result.update(sourceURL=BRENT_EIA,observations=sorted(records.values(),key=lambda r:r['period'])[-90:])
    published=re.search(r'(?<!NEXT )RELEASE DATE:\s*(\d{2}/\d{2}/\d{4})',text)
    if published:
        date=dt.datetime.strptime(published[1],'%m/%d/%Y').date()
        if date>today or date.isoformat()<result['observations'][-1]['period']:
            raise ValueError('EIA publication date invalid')
        result['publishedAt']=date.isoformat()
    elif result['observations'][-1]['period']!=previous['observations'][-1]['period']:
        result.pop('publishedAt',None)
    return result


def refresh_brent(previous, today):
    try:
        return refresh_brent_fred(previous,today)
    except Exception as primary_error:
        candidate=brent_eia_page(download(BRENT_EIA),previous,today)
        RECOVERIES['BRENT']={'primaryError':str(primary_error)[:240],'sourceURL':BRENT_EIA}
        return candidate


def refresh_data(data, today, checked):
    RECOVERIES.clear()
    result=copy.deepcopy(data)
    report={'checkedAt':checked,'changed':[],'retainedAfterError':[],'verified':[],'sources':{}}
    for code,refresh in [('CUB_GO',refresh_cub),('IPCA',refresh_ipca),('DIESEL_S10',refresh_diesel),('BRENT',refresh_brent)]:
        previous=result['series'][code]
        try:
            candidate=refresh(copy.deepcopy(previous),today)
            if candidate['observations'][-1]['period']<previous['observations'][-1]['period']:
                raise ValueError('Older reference cannot replace last verified series')
            changed=candidate!=previous
            if changed:
                candidate['verifiedAt']=checked
                result['series'][code]=candidate
                report['changed'].append(code)
            report['verified'].append(code)
            status={'status':'updated' if changed else 'verified-unchanged','period':candidate['observations'][-1]['period'],'observations':len(candidate['observations']),'sourceURL':candidate['sourceURL']}
            if code in RECOVERIES:
                status['recovery']=RECOVERIES[code]
            report['sources'][code]=status
            print(f"{code}: {status['status']}; {status['observations']} observations through {status['period']}"+(' (official publisher fallback)' if 'recovery' in status else ''))
        except Exception as error:
            detail=str(error).replace('\n',' ')[:240]
            report['retainedAfterError'].append(code)
            report['sources'][code]={'status':'retained-after-error','period':previous['observations'][-1]['period'],'error':f'{type(error).__name__}: {detail}'}
            print(f'{code}: retained last verified series ({type(error).__name__}: {detail})')
    if report['changed']:
        result['verifiedAt']=checked
    report['outcome']='unavailable' if not report['verified'] else 'partial' if report['retainedAfterError'] else 'verified'
    return result,report


def write_actions_summary(report):
    path=os.environ.get('GITHUB_STEP_SUMMARY')
    if path:
        lines=['## Published economic observations',f"Outcome: **{report['outcome']}**. Changed: {', '.join(report['changed']) or 'none'}. Source verification is separate from the observation period.",'','| Source | Status | Last reference |','|---|---|---|']
        details=[]
        for code,status in report['sources'].items():
            lines.append(f"| {code} | {status['status']} | {status['period']} |")
            if 'error' in status:
                details.append(f"{code}: {status['error'].replace('|','/')}")
            if 'recovery' in status:
                details.append(f"{code}: primary connection failed; verified official fallback {status['recovery']['sourceURL']}.")
        lines.extend(['']+details)
        with open(path,'a') as summary:
            summary.write('\n'.join(lines)+'\n')
    if os.environ.get('GITHUB_ACTIONS')=='true':
        for code,status in report['sources'].items():
            if 'error' in status:
                message=status['error'].replace('%','%25').replace('\r','%0D').replace('\n','%0A')
                print(f'::warning title={code} retained dated data::{message}')


def main():
    args = argparse.ArgumentParser()
    args.add_argument('--file', default=str(ROOT/'data/economic-series.json'))
    args.add_argument('--report')
    options = args.parse_args()
    path = Path(options.file)
    data = json.loads(path.read_text())
    today = dt.datetime.now(dt.timezone.utc).date()
    checked = dt.datetime.now(dt.timezone.utc).isoformat().replace('+00:00','Z')
    updated,report=refresh_data(data,today,checked)
    if report['changed']:
        path.write_text(json.dumps(updated,ensure_ascii=False,indent=2)+'\n')
    if options.report:
        report_path=Path(options.report);report_path.parent.mkdir(parents=True,exist_ok=True)
        report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    write_actions_summary(report)
    print(json.dumps(report))
    if report['outcome']=='unavailable':
        raise SystemExit(1)


if __name__ == '__main__':
    main()
